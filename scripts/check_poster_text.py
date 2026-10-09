#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ポスターのテキストが枠からあふれていないかを **実フォントで測って** 検査する。

【なぜ要るか】
`check_poster_layout.py` は図形の座標しか見ないので、**テキストの折り返しあふれは
原理的に検出できない**（引き継ぎメモに明記。だから「PowerPoint で開いて目視」が
残り作業になっていた）。この検査はそこを埋める:

  * pptx の各テキストボックスについて、`scripts/textmetrics.py` で実際の Arial を
    読み、run ごとの字幅を測って**本当の折り返し行数**を求める
  * 行高は Arial の実メトリクス（ascent+descent）× 段落の line spacing
  * 必要高がボックス高を超えたものを列・パネル単位で並べる

**組む側（build_poster_pptx.py の est_h）と同じ `textmetrics` で測る。**
別々に見積もると片方だけ直したときに検算が素通りする。

    python3 scripts/check_poster_text.py [pptx]   # 既定は docs/poster/P202_poster.pptx

【読み方】
ボックス高は build 側が「入るはず」と見込んだ高さなので、超過 = 見込み違い。
ただし **超過がすぐ事故とは限らない**：下に余白がある枠なら文字が下へ伸びるだけで
実害が無い。そこで「超過ぶんが、その枠の下にある次の図形までの空きを食うか」まで見て、
**衝突するものだけを 🔴** にする。

    python3 scripts/check_poster_text.py
    → 🔴 が 0 なら、刷ってよい（それでも一度は実物を見ること）
"""
import sys, pathlib, json
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import textmetrics as TM

ROOT = pathlib.Path(__file__).resolve().parent.parent
PPTX = pathlib.Path(sys.argv[1]).resolve() if len(sys.argv) > 1 \
    else ROOT / "docs/poster/P202_poster.pptx"
LAYOUT = ROOT / "docs/poster/layout.json"

if not TM.available():
    sys.exit("🔴 Arial が読めない。実測できないので検査を中止する "
             "（近似で通すと検査の意味が無い）")


def main():
    if not PPTX.exists():
        sys.exit(f"🔴 {PPTX} が無い。先に build_poster_pptx.py を流すこと")
    L = json.loads(LAYOUT.read_text())
    SCALE = L["SCALE"]
    prs = Presentation(str(PPTX))
    slide = prs.slides[0]

    def col_of(x_in):
        for i in range(len(L["X"]) - 1, -1, -1):
            if x_in >= L["X"][i] - 0.01:
                return i + 1
        return 0

    # 刷り上がり寸法（in）に直して集める。
    # 🔴 グループの中まで降りること。pptx を PowerPoint で手編集すると図形が
    # グループ化されることがあり、`slide.shapes` はグループを 1 個としか数えない。
    # 2026-09-21: 9/10 版の 488 図形が 225 としか見えず、中の 263 個が
    # **検査されないまま素通りしていた**のに気づいた。
    from pptx.oxml.ns import qn

    def children(shape):
        """グループの子と、子の座標を絶対座標へ直すための変換を返す。"""
        x = shape._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
        off, ext = x.find(qn("a:off")), x.find(qn("a:ext"))
        coff, cext = x.find(qn("a:chOff")), x.find(qn("a:chExt"))
        sx = int(ext.get("cx")) / max(1, int(cext.get("cx")))
        sy = int(ext.get("cy")) / max(1, int(cext.get("cy")))
        return (int(off.get("x")), int(off.get("y")),
                int(coff.get("x")), int(coff.get("y")), sx, sy)

    boxes, solids = [], []

    def walk(shapes, tf_=(0, 0, 0, 0, 1.0, 1.0)):
        ox, oy, cx0, cy0, sx, sy = tf_
        for sh in shapes:
            left = ox + (sh.left - cx0) * sx
            top = oy + (sh.top - cy0) * sy
            wid, hgt = sh.width * sx, sh.height * sy
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                gx, gy, gcx, gcy, gsx, gsy = children(sh)
                walk(sh.shapes, (ox + (gx - cx0) * sx, oy + (gy - cy0) * sy,
                                 gcx, gcy, sx * gsx, sy * gsy))
                continue
            x, y = Emu(int(left)).inches / SCALE, Emu(int(top)).inches / SCALE
            w, h = Emu(int(wid)).inches / SCALE, Emu(int(hgt)).inches / SCALE
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                need = 0.0
                for p_ in tf.paragraphs:
                    runs = [(r.text, (r.font.size.pt if r.font.size else 12) / SCALE,
                             r.font.bold, r.font.italic) for r in p_.runs]
                    if not runs:
                        continue
                    sp = p_.line_spacing if isinstance(p_.line_spacing, float) else 1.0
                    before = (p_.space_before.pt / SCALE) if p_.space_before else 0.0
                    sizes = ([max(r[1] for r in runs)] if tf.word_wrap is False
                             else TM.wrap(runs, w * 72))
                    need += before / 72 + sum(TM.line_h_pt(s) * sp for s in sizes) / 72
                boxes.append(dict(x=x, y=y, w=w, h=h, need=need,
                                  txt=tf.text.strip().replace("\n", " / ")[:78]))
            else:
                solids.append((x, y, w, h))

    walk(slide.shapes)

    # 「下にある次の図形までの空き」を測る。横に重なるものだけを相手にする。
    for b in boxes:
        below = [yy for (xx, yy, ww, hh) in solids + [(o["x"], o["y"], o["w"], o["h"])
                                                      for o in boxes if o is not b]
                 if yy >= b["y"] + b["h"] - 0.02
                 and xx < b["x"] + b["w"] - 0.02 and xx + ww > b["x"] + 0.02]
        b["room"] = (min(below) - (b["y"] + b["h"])) if below else 99.0
        b["over"] = b["need"] - b["h"]

    over = sorted([b for b in boxes if b["over"] > 0.02],
                  key=lambda b: -(b["over"] - b["room"]))
    hard = [b for b in over if b["over"] > b["room"] + 0.02]

    print(f"{PPTX.name}: テキストボックス {len(boxes)} 個を実フォント（Arial）で測った\n")
    if not over:
        print("✅ 枠からあふれる文字は無い")
    else:
        print(f"見込みより高くなった枠 {len(over)} 個 "
              f"（🔴 = 下の図形に当たる / △ = 下の空きに収まる）\n")
        for b in over:
            mark = "🔴" if b["over"] > b["room"] + 0.02 else "△"
            room = "—" if b["room"] > 90 else f"{b['room']:.2f}"
            print(f" {mark} 列{col_of(b['x'])} y={b['y']:6.2f} "
                  f"枠{b['h']:5.2f} 要{b['need']:5.2f} "
                  f"超過{b['over']:+5.2f} 下の空き {room}")
            print(f"      {b['txt']}")
    print(f"\n🔴 {len(hard)} 件 / △ {len(over) - len(hard)} 件")
    print("※ この検査は折り返しと行高を実測する。色・重なり・図の中身は見ていない。")
    return 1 if hard else 0


if __name__ == "__main__":
    sys.exit(main())
